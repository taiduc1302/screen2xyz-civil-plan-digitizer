"""Estimator XLSX/point profiles and strict generated-file round-trip checks."""

from __future__ import annotations

import csv
import io
import json
import os
from pathlib import Path
from typing import Any, Iterable

from . import PRELIMINARY_WARNING
from . import contracts as C
from .io_utils import atomic_write_bytes
from .models import CivilPoint, CivilProject
from .transform import to_local


POINT_COLUMNS = (
    "Point",
    "Northing",
    "Easting",
    "Elevation",
    "Description",
    "Class",
    "Page",
    "Sheet",
    "Revision",
    "Coordinate System",
    "Source Method",
    "Confidence",
    "Review Status",
    "Notes",
    "Created Timestamp",
    "Last Edited Timestamp",
)

AGTEK_COLUMNS = ("Point", "Northing", "Easting", "Elevation", "Description")
XYZ_COLUMNS = ("X", "Y", "Z", "Description")
NEZ_COLUMNS = AGTEK_COLUMNS


class EstimatorExportError(RuntimeError):
    """Estimator output could not be created or verified safely."""


def point_number(point: CivilPoint) -> str:
    return point.point_number or point.id


def generic_xyz_bytes(
    points: Iterable[CivilPoint],
    *,
    precision: int = 6,
    delimiter: str = ",",
) -> bytes:
    rows = [
        {
            "X": _number(point.local_east, precision),
            "Y": _number(point.local_north, precision),
            "Z": _number(point.elevation, precision),
            "Description": _safe_text(point.description),
        }
        for point in points
    ]
    return _csv_bytes(rows, XYZ_COLUMNS, delimiter=delimiter)


def generic_nez_bytes(
    points: Iterable[CivilPoint],
    *,
    precision: int = 6,
    delimiter: str = ",",
) -> bytes:
    rows = [
        {
            "Point": point_number(point),
            "Northing": _number(point.local_north, precision),
            "Easting": _number(point.local_east, precision),
            "Elevation": _number(point.elevation, precision),
            "Description": _safe_text(point.description),
        }
        for point in points
    ]
    return _csv_bytes(rows, NEZ_COLUMNS, delimiter=delimiter)


def write_estimator_workbook(
    project: CivilProject,
    path: Path,
    points: Iterable[CivilPoint],
    *,
    qa_summary: dict[str, Any],
    coordinate_system: str = "LOCAL EAST/NORTH (metres; non-geodetic)",
) -> Path:
    """Write a practical eight-sheet workbook using the pinned local library."""

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise EstimatorExportError(
            "XLSX export requires openpyxl from requirements.txt"
        ) from exc

    approved = list(points)
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = f"{project.name} approved civil points"
    workbook.properties.subject = "Preliminary estimator point export"
    workbook.properties.creator = "Screen2XYZ Civil Plan Digitizer"
    workbook.properties.description = PRELIMINARY_WARNING

    thin = Side(style="thin", color="D8DEE9")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    warning_fill = PatternFill("solid", fgColor="FCE8E6")

    def add_point_sheet(name: str, selected: list[CivilPoint]) -> None:
        sheet = workbook.create_sheet(name)
        sheet.sheet_view.showGridLines = False
        sheet.append(list(POINT_COLUMNS))
        for point in selected:
            sheet.append(
                _workbook_point_row(
                    point,
                    coordinate_system=coordinate_system,
                )
            )
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:P{max(1, sheet.max_row)}"
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=thin)
        for row in sheet.iter_rows(min_row=2):
            for index in (2, 3, 4, 12):
                row[index - 1].number_format = "0.000000"
            for cell in row:
                cell.alignment = Alignment(vertical="top")
        widths = (
            14,
            15,
            15,
            13,
            22,
            20,
            8,
            10,
            10,
            32,
            16,
            12,
            20,
            28,
            22,
            22,
        )
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.row_dimensions[1].height = 28

    existing = [
        point for point in approved if point.point_type == C.EXISTING_GROUND
    ]
    design = [point for point in approved if point.point_type == C.DESIGN_GRADE]
    contours = [
        point for point in approved if point.point_type == C.CONTOUR_ELEVATION
    ]
    add_point_sheet("Existing Points", existing)
    add_point_sheet("Design Points", design)
    add_point_sheet("Contours and Lines", contours)
    contour_sheet = workbook["Contours and Lines"]
    for line in project.elevation_lines:
        if not line.approved or project.calibration is None:
            continue
        for vertex_index, vertex in enumerate(line.vertices, start=1):
            east, north = to_local(vertex, project.calibration)
            contour_sheet.append(
                [
                    _safe_text(f"{line.id}-V{vertex_index:03d}"),
                    north,
                    east,
                    line.elevation,
                    _safe_text(line.description),
                    "CONTOUR_LINE_VERTEX",
                    project.source_manifest.get("selected_page_index", 0) + 1,
                    _safe_text(line.sheet),
                    _safe_text(line.revision_label),
                    _safe_text(coordinate_system),
                    _safe_text(line.source_method),
                    None,
                    _safe_text(line.review_status),
                    _safe_text(line.notes),
                    _safe_text(line.created_at),
                    _safe_text(line.updated_at),
                ]
            )
    contour_sheet.auto_filter.ref = f"A1:P{max(1, contour_sheet.max_row)}"
    for row in contour_sheet.iter_rows(min_row=2):
        for index in (2, 3, 4, 12):
            row[index - 1].number_format = "0.000000"
        for cell in row:
            cell.alignment = Alignment(vertical="top")
    add_point_sheet("All Approved Data", approved)

    qa_sheet = workbook.create_sheet("QA Summary")
    qa_sheet.sheet_view.showGridLines = False
    qa_sheet.append(["Check", "Value"])
    qa_sheet.append(["Preliminary warning", PRELIMINARY_WARNING])
    for key in sorted(qa_summary):
        value = qa_summary[key]
        qa_sheet.append(
            [
                _safe_text(key),
                (
                    json.dumps(value, ensure_ascii=False, sort_keys=True)
                    if isinstance(value, (dict, list))
                    else (
                        _safe_text(value)
                        if isinstance(value, str)
                        else value
                    )
                ),
            ]
        )
    _style_key_value_sheet(qa_sheet, header_fill, header_font, warning_fill)

    calibration_sheet = workbook.create_sheet("Calibration")
    calibration_sheet.sheet_view.showGridLines = False
    calibration_sheet.append(["Field", "Value"])
    calibration_sheet.append(["Preliminary warning", PRELIMINARY_WARNING])
    if project.calibration is None:
        calibration_sheet.append(["Status", "NOT CALIBRATED"])
    else:
        calibration = project.calibration
        rows = (
            ("Revision", calibration.revision),
            ("Metres per pixel", calibration.metres_per_pixel),
            ("Origin pixel X", calibration.origin_pixel.x),
            ("Origin pixel Y", calibration.origin_pixel.y),
            ("Origin Easting", calibration.origin_east_m),
            ("Origin Northing", calibration.origin_north_m),
            ("East unit X", calibration.east_unit_x),
            ("East unit Y", calibration.east_unit_y),
            ("Coordinate system", coordinate_system),
        )
        for row in rows:
            calibration_sheet.append(
                [
                    _safe_text(row[0]),
                    (
                        _safe_text(row[1])
                        if isinstance(row[1], str)
                        else row[1]
                    ),
                ]
            )
    _style_key_value_sheet(
        calibration_sheet, header_fill, header_font, warning_fill
    )
    for cell in calibration_sheet["B"][2:]:
        if isinstance(cell.value, (int, float)):
            cell.number_format = "0.000000000000"

    source_sheet = workbook.create_sheet("Source Metadata")
    source_sheet.sheet_view.showGridLines = False
    source_sheet.append(["Field", "Value"])
    allowed = {
        "display_name",
        "sha256",
        "byte_size",
        "width_px",
        "height_px",
        "source_type",
        "page_count",
        "selected_page_index",
        "sheet_id",
        "revision",
    }
    for key in sorted(allowed):
        if key in project.source_manifest:
            value = project.source_manifest[key]
            source_sheet.append(
                [
                    _safe_text(key),
                    _safe_text(value) if isinstance(value, str) else value,
                ]
            )
    source_sheet.append(["Local path exported", False])
    _style_key_value_sheet(source_sheet, header_fill, header_font, warning_fill)

    settings_sheet = workbook.create_sheet("Export Settings")
    settings_sheet.sheet_view.showGridLines = False
    settings_sheet.append(["Setting", "Value"])
    for row in (
        ("Default policy", "Approved terrain points only"),
        ("Rejected included", False),
        ("Unreviewed included", False),
        ("Coordinate system", coordinate_system),
        ("Coordinate units", "metres"),
        ("Coordinate precision", 6),
        ("Elevation precision", 6),
        ("Existing/Design separation", True),
        ("AGTEK compatibility", "PREPARED; downstream import not yet verified"),
        ("Kubla compatibility", "CONFIGURABLE; official/real import not yet verified"),
    ):
        settings_sheet.append(
            [
                _safe_text(row[0]),
                _safe_text(row[1]) if isinstance(row[1], str) else row[1],
            ]
        )
    _style_key_value_sheet(settings_sheet, header_fill, header_font, warning_fill)

    target = path.expanduser().resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = target.with_name(f".{target.name}.tmp")
    try:
        workbook.save(temporary)
        os.replace(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)
    return target


def verify_export_round_trip(
    project: CivilProject,
    export_dir: Path,
    expected_points: Iterable[CivilPoint],
    *,
    coordinate_order: str,
    precision: int = 6,
) -> dict[str, Any]:
    """Reopen every required point file and compare it to the approved cart."""

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise EstimatorExportError(
            "round-trip verification requires openpyxl==3.1.5"
        ) from exc

    expected = list(expected_points)
    expected_by_number = {point_number(point): point for point in expected}
    checks: list[dict[str, Any]] = []

    def record(name: str, passed: bool, detail: str) -> None:
        checks.append({"name": name, "passed": bool(passed), "detail": detail})

    if len(expected_by_number) != len(expected):
        record("unique approved point numbers", False, "duplicate point numbers")
    else:
        record("unique approved point numbers", True, f"{len(expected)} unique")

    workbook_path = export_dir / "Approved_Point_Cart.xlsx"
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        all_sheet = workbook["All Approved Data"]
        headers = [cell.value for cell in next(all_sheet.iter_rows(min_row=1, max_row=1))]
        rows = [
            dict(zip(headers, values))
            for values in all_sheet.iter_rows(min_row=2, values_only=True)
            if any(value is not None for value in values)
        ]
        _verify_point_rows(
            rows,
            expected_by_number,
            record=record,
            label="XLSX All Approved Data",
            precision=precision,
        )
        split_counts = {
            "Existing Points": sum(
                point.point_type == C.EXISTING_GROUND for point in expected
            ),
            "Design Points": sum(
                point.point_type == C.DESIGN_GRADE for point in expected
            ),
            "Contours and Lines": sum(
                point.point_type == C.CONTOUR_ELEVATION for point in expected
            )
            + sum(
                len(line.vertices)
                for line in project.elevation_lines
                if line.approved
            ),
        }
        for sheet_name, count in split_counts.items():
            sheet = workbook[sheet_name]
            actual = sum(
                1
                for row in sheet.iter_rows(min_row=2, values_only=True)
                if any(value is not None for value in row)
            )
            record(
                f"XLSX {sheet_name} separation",
                actual == count,
                f"expected {count}; reopened {actual}",
            )
    finally:
        workbook.close()

    all_csv_rows = _read_csv(export_dir / "All_Reviewed_Points.csv")
    _verify_agtek_rows(
        all_csv_rows,
        expected_by_number,
        record=record,
        label="All reviewed CSV",
        coordinate_order=coordinate_order,
        precision=precision,
    )
    for name, point_type in (
        ("Existing_Points.csv", C.EXISTING_GROUND),
        ("Design_Points.csv", C.DESIGN_GRADE),
    ):
        rows = _read_csv(export_dir / name)
        expected_subset = {
            number: point
            for number, point in expected_by_number.items()
            if point.point_type == point_type
        }
        _verify_agtek_rows(
            rows,
            expected_subset,
            record=record,
            label=name,
            coordinate_order=coordinate_order,
            precision=precision,
        )

    xyz_rows = _read_csv(export_dir / "Reviewed_Points.xyz")
    expected_ordered = list(expected)
    xyz_ok = len(xyz_rows) == len(expected_ordered)
    if xyz_ok:
        for row, point in zip(xyz_rows, expected_ordered):
            xyz_ok = xyz_ok and _equal_number(
                row.get("X"), point.local_east, precision
            )
            xyz_ok = xyz_ok and _equal_number(
                row.get("Y"), point.local_north, precision
            )
            xyz_ok = xyz_ok and _equal_number(
                row.get("Z"), point.elevation, precision
            )
            xyz_ok = xyz_ok and row.get("Description", "") == _safe_text(
                point.description
            )
    record(
        "generic XYZ round trip",
        xyz_ok,
        f"expected {len(expected_ordered)}; reopened {len(xyz_rows)}",
    )

    nez_rows = _read_csv(export_dir / "Reviewed_Points.nez")
    _verify_agtek_rows(
        nez_rows,
        expected_by_number,
        record=record,
        label="generic NEZ",
        coordinate_order="NE",
        precision=precision,
    )

    safety_ok = all(
        point.approved
        and point.terrain_relevant
        and "%" not in str(point.elevation)
        and point.local_east is not None
        and point.local_north is not None
        for point in expected
    )
    record(
        "approved-only safety policy",
        safety_ok,
        "no rejected, pending, non-terrain, percent, or empty-coordinate row",
    )
    passed = all(check["passed"] for check in checks)
    report = _round_trip_report(checks, passed, len(expected))
    atomic_write_bytes(
        export_dir / "Export_Round_Trip_Report.md",
        report.encode("utf-8"),
        replace=True,
    )
    return {
        "passed": passed,
        "checks": checks,
        "approved_count": len(expected),
        "report_path": export_dir / "Export_Round_Trip_Report.md",
    }


def _workbook_point_row(
    point: CivilPoint,
    *,
    coordinate_system: str,
) -> list[Any]:
    return [
        _safe_text(point_number(point)),
        float(point.local_north),
        float(point.local_east),
        float(point.elevation),
        _safe_text(point.description),
        _safe_text(point.point_type),
        point.page_index + 1,
        _safe_text(point.sheet or point.page_label),
        _safe_text(point.revision_label),
        _safe_text(coordinate_system),
        _safe_text(point.source_method),
        (
            point.capture_confidence
            if point.capture_confidence is not None
            else point.classification_confidence
        ),
        _safe_text(point.review_status),
        _safe_text(point.notes),
        _safe_text(point.created_at),
        _safe_text(point.updated_at),
    ]


def _style_key_value_sheet(sheet, header_fill, header_font, warning_fill) -> None:
    from openpyxl.styles import Alignment, Border, Side

    thin = Side(style="thin", color="D8DEE9")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = Border(bottom=thin)
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 90
    sheet["B2"].fill = warning_fill
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _verify_point_rows(
    rows: list[dict[str, Any]],
    expected_by_number: dict[str, CivilPoint],
    *,
    record,
    label: str,
    precision: int,
) -> None:
    numbers = [str(row.get("Point", "")) for row in rows]
    ids_ok = len(numbers) == len(set(numbers)) and set(numbers) == set(
        expected_by_number
    )
    record(
        f"{label} IDs/count",
        ids_ok,
        f"expected {len(expected_by_number)}; reopened {len(rows)}",
    )
    values_ok = ids_ok
    if values_ok:
        for row in rows:
            point = expected_by_number[str(row["Point"])]
            values_ok = values_ok and _equal_number(
                row.get("Northing"), point.local_north, precision
            )
            values_ok = values_ok and _equal_number(
                row.get("Easting"), point.local_east, precision
            )
            values_ok = values_ok and _equal_number(
                row.get("Elevation"), point.elevation, precision
            )
            values_ok = values_ok and row.get("Class") == point.point_type
            values_ok = values_ok and row.get("Description", "") == _safe_text(
                point.description
            )
            values_ok = values_ok and row.get("Review Status") in C.APPROVED_STATUSES
    record(
        f"{label} numeric/class/description values",
        values_ok,
        "numeric cells, class, description, and approval state compared",
    )


def _verify_agtek_rows(
    rows: list[dict[str, str]],
    expected_by_number: dict[str, CivilPoint],
    *,
    record,
    label: str,
    coordinate_order: str,
    precision: int,
) -> None:
    numbers = [row.get("Point", "") for row in rows]
    ids_ok = len(numbers) == len(set(numbers)) and set(numbers) == set(
        expected_by_number
    )
    values_ok = ids_ok
    if values_ok:
        for row in rows:
            point = expected_by_number[row["Point"]]
            expected_north = (
                point.local_north if coordinate_order == "NE" else point.local_east
            )
            expected_east = (
                point.local_east if coordinate_order == "NE" else point.local_north
            )
            values_ok = values_ok and _equal_number(
                row.get("Northing"), expected_north, precision
            )
            values_ok = values_ok and _equal_number(
                row.get("Easting"), expected_east, precision
            )
            values_ok = values_ok and _equal_number(
                row.get("Elevation"), point.elevation, precision
            )
            values_ok = values_ok and row.get("Description", "") == _safe_text(
                point.description
            )
    record(
        f"{label} round trip",
        values_ok,
        f"expected {len(expected_by_number)}; reopened {len(rows)}",
    )


def _round_trip_report(
    checks: list[dict[str, Any]],
    passed: bool,
    approved_count: int,
) -> str:
    lines = [
        "# Export Round-Trip Report",
        "",
        f"> {PRELIMINARY_WARNING}",
        "",
        f"- Overall: {'PASS' if passed else 'FAIL'}",
        f"- Approved Point Cart rows expected: {approved_count}",
        "",
        "## Checks",
        "",
    ]
    lines.extend(
        f"- [{'PASS' if check['passed'] else 'FAIL'}] "
        f"{check['name']}: {check['detail']}"
        for check in checks
    )
    return "\n".join(lines) + "\n"


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def _csv_bytes(
    rows: list[dict[str, str]],
    headers: tuple[str, ...],
    *,
    delimiter: str,
) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.DictWriter(
        buffer,
        fieldnames=list(headers),
        delimiter=delimiter,
        lineterminator="\r\n",
        extrasaction="ignore",
    )
    writer.writeheader()
    writer.writerows(rows)
    return ("\ufeff" + buffer.getvalue()).encode("utf-8")


def _number(value: float | None, precision: int) -> str:
    if value is None:
        raise EstimatorExportError("approved export point has an empty coordinate")
    return format(float(value), f".{precision}f")


def _equal_number(actual: Any, expected: float | None, precision: int) -> bool:
    if actual is None or expected is None:
        return False
    try:
        return round(float(actual), precision) == round(float(expected), precision)
    except (TypeError, ValueError):
        return False


def _safe_text(value: Any) -> str:
    text = str(value or "")
    return f"'{text}" if text.startswith(("=", "+", "-", "@")) else text
