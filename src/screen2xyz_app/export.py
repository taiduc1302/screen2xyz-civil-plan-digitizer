"""Default XLSX/CSV exports and the opt-in advanced estimator handoff."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Iterable

from screen2xyz_civil.estimator_exports import (
    _safe_text,
    _style_key_value_sheet,
)
from screen2xyz_civil.exports import export_handoff
from screen2xyz_civil import contracts as civil_contracts
from screen2xyz_civil.models import CivilPoint, CivilProject

from .store import SessionStore

POINT_HEADERS = (
    "PointNumber", "X", "Y", "Z", "Status", "Description", "SourceX", "SourceY",
    "SourceZ", "Confidence", "Timestamp",
)
PRELIMINARY_WARNING = (
    "Preliminary data for conceptual estimating only; validate against an "
    "authoritative source before use. Not certified survey data."
)


def _rows(store: SessionStore, session_id: str) -> Iterable[list[object]]:
    for sequence, point in enumerate(store.points(session_id), start=1):
        confidences = [
            point[key] for key in ("confidence_x", "confidence_y", "confidence_z")
            if point[key] is not None
        ]
        yield [
            _safe_text(point["point_number"] or str(sequence)),
            point["x"], point["y"], point["z"], point["capture_status"],
            _safe_text(point["description"]),
            _safe_text(point["source_method_x"]),
            _safe_text(point["source_method_y"]),
            _safe_text(point["source_method_z"]),
            None if not confidences else sum(confidences) / len(confidences),
            _safe_text(point["created_utc"]),
        ]


def export_xlsx(store: SessionStore, session_id: str, output_path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    session = store.session(session_id)
    workbook = Workbook()
    points = workbook.active
    points.title = "Points"
    points.append(POINT_HEADERS)
    for row in _rows(store, session_id):
        points.append(row)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    warning_fill = PatternFill("solid", fgColor="FFF2CC")
    for cell in points[1]:
        cell.fill = header_fill
        cell.font = header_font
    points.freeze_panes = "A2"
    points.auto_filter.ref = points.dimensions
    for index, width in enumerate(
        (14, 16, 16, 14, 24, 34, 20, 20, 20, 14, 28), start=1
    ):
        points.column_dimensions[get_column_letter(index)].width = width
    for row in points.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    details = workbook.create_sheet("Session")
    details.append(("Field", "Value"))
    details.append(("Warning", PRELIMINARY_WARNING))
    details.append(("Session ID", session["id"]))
    details.append(("Created UTC", session["created_utc"]))
    details.append(("App version", session["app_version"]))
    details.append(("Channel mapping", json.dumps(json.loads(session["mapping_json"]), indent=2)))
    details.append(("Calibration", session["calibration_json"] or "None"))
    _style_key_value_sheet(details, header_fill, header_font, warning_fill)

    output_path = output_path.expanduser().resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_suffix(output_path.suffix + ".tmp")
    workbook.save(temporary)
    temporary.replace(output_path)
    return output_path


def export_csv(store: SessionStore, session_id: str, output_path: Path) -> Path:
    output_path = output_path.expanduser().resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_suffix(output_path.suffix + ".tmp")
    with temporary.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\r\n")
        writer.writerow(POINT_HEADERS)
        writer.writerows(_rows(store, session_id))
    temporary.replace(output_path)
    return output_path


def advanced_estimator_export(
    store: SessionStore, session_id: str, output_root: Path
) -> dict[str, object]:
    """Build the retained multi-artifact Civil estimator handoff."""
    session = store.session(session_id)
    points = []
    for sequence, row in enumerate(store.points(session_id), start=1):
        if row["z"] is None:
            raise ValueError(
                "advanced estimator export requires complete rows; "
                "review PARTIAL_MISSING_Z rows first"
            )
        points.append(CivilPoint(
            id=f"PT-{sequence:06d}", page_index=0, page_label="1",
            source_file="Screen2XYZ v2 session", source_sha256="",
            pixel_x=0.0, pixel_y=0.0, elevation=float(row["z"]),
            local_east=float(row["x"]), local_north=float(row["y"]),
            point_type=civil_contracts.EXISTING_GROUND, symbol_type="MANUAL_POINT",
            source_method=civil_contracts.MANUAL,
            review_status=civil_contracts.APPROVED,
            created_at=row["created_utc"], updated_at=row["created_utc"],
            description=row["description"] or "",
            point_number=row["point_number"] or str(sequence),
        ))
    project = CivilProject(
        project_id=f"CPD-V2-{session_id}", name="Screen2XYZ-v2",
        created_at=session["created_utc"], updated_at=session["created_utc"],
        source_manifest={"display_name": "Screen2XYZ v2 session"},
        points=points,
        feature_flags={
            "preliminary_surface": False,
            "landxml": False,
            "external_coordinate_channels": True,
        },
    )
    return export_handoff(
        project, output_root,
        now=lambda: session["created_utc"],
    )
