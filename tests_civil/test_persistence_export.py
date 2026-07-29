from __future__ import annotations

import csv
import json
import unittest
from unittest import mock

from screen2xyz_civil import PRELIMINARY_WARNING
from screen2xyz_civil import contracts as C
from screen2xyz_civil.estimator_exports import verify_export_round_trip
from screen2xyz_civil.exports import (
    AGTEK_HEADER,
    ExportError,
    approved_points,
    export_handoff,
)
from screen2xyz_civil.models import PixelPoint
from screen2xyz_civil.persistence import (
    ProjectPersistenceError,
    load_project,
    migrate_project_dict,
    save_project,
)

from .helpers_civil import add_approved, clock, fresh_dir, project_and_workflow


class PersistenceExportTests(unittest.TestCase):
    def test_approved_contour_line_exports_geometry_and_vertices(self):
        root = fresh_dir()
        project, flow = project_and_workflow()
        add_approved(flow)
        line = flow.add_elevation_line(
            [PixelPoint(60, 90), PixelPoint(90, 80)],
            elevation=49.5,
        )
        flow.approve_elevation_line(line.id)
        saved = root / "contours.s2c.json"
        save_project(project, saved)
        self.assertEqual(load_project(saved).elevation_lines[0].id, line.id)
        result = export_handoff(project, root, now=clock, export_id="contours")
        geometry = json.loads(
            (result["export_dir"] / "Contour_Lines.geojson").read_text(
                encoding="utf-8"
            )
        )
        self.assertEqual(geometry["features"][0]["id"], line.id)
        with (result["export_dir"] / "Contour_Line_Vertices.csv").open(
            encoding="utf-8-sig", newline=""
        ) as handle:
            rows = list(csv.DictReader(handle))
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["Elevation"], "49.500000")

    def test_project_save_load_round_trip(self):
        root = fresh_dir()
        project, flow = project_and_workflow()
        add_approved(flow)
        path = root / "project.s2c.json"
        identity = save_project(project, path)
        restored = load_project(path)
        self.assertEqual(restored.to_dict(), project.to_dict())
        self.assertEqual(len(identity["sha256"]), 64)

    def test_wrong_project_suffix_rejected(self):
        project, _flow = project_and_workflow()
        with self.assertRaises(ProjectPersistenceError):
            save_project(project, fresh_dir() / "project.json")

    def test_schema_09_migrates(self):
        project, _flow = project_and_workflow()
        value = project.to_dict()
        value["schema_version"] = "0.9"
        value["points"] = value.pop("point_candidates")
        migrated = migrate_project_dict(value)
        self.assertEqual(migrated["schema_version"], C.SCHEMA_VERSION)
        self.assertIn("point_candidates", migrated)

    def test_future_schema_rejected(self):
        with self.assertRaises(ProjectPersistenceError):
            migrate_project_dict({"schema_version": "99.0"})

    def test_approved_only_existing_design_split(self):
        root = fresh_dir()
        project, flow = project_and_workflow()
        add_approved(flow, point_type=C.EXISTING_GROUND)
        add_approved(
            flow,
            pixel=PixelPoint(80, 80),
            elevation=50,
            point_type=C.DESIGN_GRADE,
        )
        flow.add_manual_point(
            pixel=PixelPoint(90, 70),
            elevation=51,
            point_type=C.DESIGN_GRADE,
        )
        result = export_handoff(
            project, root, now=clock, export_id="deterministic"
        )
        self.assertEqual(result["approved_count"], 2)
        self.assertEqual(result["existing_count"], 1)
        self.assertEqual(result["design_count"], 1)
        with (result["export_dir"] / "Design_Points.csv").open(
            encoding="utf-8-sig", newline=""
        ) as handle:
            rows = list(csv.DictReader(handle))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["Description"], "MANUAL_DES")

    def test_agtek_column_order(self):
        root = fresh_dir()
        project, flow = project_and_workflow()
        add_approved(flow)
        result = export_handoff(project, root, now=clock, export_id="columns")
        first_line = (
            result["export_dir"] / "Existing_Points.csv"
        ).read_text(encoding="utf-8-sig").splitlines()[0]
        self.assertEqual(tuple(first_line.split(",")), AGTEK_HEADER)

    def test_export_contains_audit_warning_and_required_files(self):
        root = fresh_dir()
        project, flow = project_and_workflow()
        add_approved(flow)
        result = export_handoff(project, root, now=clock, export_id="package")
        required = {
            "Existing_Points.csv",
            "Design_Points.csv",
            "All_Reviewed_Points.csv",
            "Project_Audit.json",
            "Calibration_Report.md",
            "QA_Summary.md",
            "AGTEK_Import_Instructions.md",
            "Source_Manifest.json",
            "Handoff_Manifest.json",
            "evidence_manifest_sha256.txt",
            "Reviewed_Points.xyz",
            "Reviewed_Points.nez",
            "Reviewed_Points.geojson",
            "Reviewed_Points.dxf",
            "Breaklines.geojson",
            "Contour_Lines.geojson",
            "Contour_Line_Vertices.csv",
            "Approved_Point_Cart.xlsx",
            "Export_Round_Trip_Report.md",
        }
        self.assertEqual(
            required,
            {path.name for path in result["export_dir"].iterdir() if path.is_file()},
        )
        audit = json.loads(
            (result["export_dir"] / "Project_Audit.json").read_text(
                encoding="utf-8"
            )
        )
        self.assertEqual(audit["warning"], PRELIMINARY_WARNING)

    def test_xlsx_has_required_sheets_numeric_cells_and_filters(self):
        from openpyxl import load_workbook

        root = fresh_dir()
        project, flow = project_and_workflow()
        add_approved(flow)
        project.source_manifest["display_name"] = "=unsafe-formula.png"
        line = flow.add_elevation_line(
            [PixelPoint(60, 90), PixelPoint(90, 80)],
            elevation=49.5,
        )
        flow.approve_elevation_line(line.id)
        result = export_handoff(project, root, now=clock, export_id="xlsx")
        workbook = load_workbook(
            result["export_dir"] / "Approved_Point_Cart.xlsx",
            read_only=False,
            data_only=True,
        )
        try:
            self.assertEqual(
                workbook.sheetnames,
                [
                    "Existing Points",
                    "Design Points",
                    "Contours and Lines",
                    "All Approved Data",
                    "QA Summary",
                    "Calibration",
                    "Source Metadata",
                    "Export Settings",
                ],
            )
            sheet = workbook["All Approved Data"]
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertTrue(sheet.auto_filter.ref)
            self.assertIsInstance(sheet["B2"].value, (int, float))
            self.assertIsInstance(sheet["C2"].value, (int, float))
            self.assertIsInstance(sheet["D2"].value, (int, float))
            contour_sheet = workbook["Contours and Lines"]
            self.assertEqual(contour_sheet.auto_filter.ref, "A1:P3")
            self.assertIsInstance(contour_sheet["B2"].value, (int, float))
            source_rows = {
                row[0].value: row[1].value
                for row in workbook["Source Metadata"].iter_rows()
            }
            self.assertEqual(
                source_rows["display_name"],
                "'=unsafe-formula.png",
            )
        finally:
            workbook.close()
        report = (
            result["export_dir"] / "Export_Round_Trip_Report.md"
        ).read_text(encoding="utf-8")
        self.assertIn("- Overall: PASS", report)

    def test_round_trip_detects_modified_point_file(self):
        root = fresh_dir()
        project, flow = project_and_workflow()
        add_approved(flow)
        result = export_handoff(project, root, now=clock, export_id="tamper")
        existing_path = result["export_dir"] / "Existing_Points.csv"
        text = existing_path.read_text(encoding="utf-8-sig")
        existing_path.write_text(
            text.replace("49.060000", "99.990000"),
            encoding="utf-8-sig",
            newline="",
        )
        verification = verify_export_round_trip(
            project,
            result["export_dir"],
            approved_points(project),
            coordinate_order="NE",
        )
        self.assertFalse(verification["passed"])
        self.assertTrue(
            any(
                not check["passed"] and check["name"] == "Existing_Points.csv round trip"
                for check in verification["checks"]
            )
        )

    def test_source_manifest_redacts_local_path(self):
        root = fresh_dir()
        project, flow = project_and_workflow()
        add_approved(flow)
        result = export_handoff(project, root, now=clock, export_id="privacy")
        text = (result["export_dir"] / "Source_Manifest.json").read_text(
            encoding="utf-8"
        )
        self.assertNotIn("private/local", text)
        self.assertIn('"local_path_exported":false', text)

    def test_export_folder_is_copy_on_write(self):
        root = fresh_dir()
        project, flow = project_and_workflow()
        add_approved(flow)
        export_handoff(project, root, now=clock, export_id="same")
        with self.assertRaises(ExportError):
            export_handoff(project, root, now=clock, export_id="same")

    def test_failed_export_leaves_no_partial_final_or_staging_folder(self):
        root = fresh_dir()
        project, flow = project_and_workflow()
        add_approved(flow)
        prior_history = list(project.export_history)
        prior_stale = project.exports_stale
        with mock.patch(
            "screen2xyz_civil.exports.verify_export_round_trip",
            side_effect=OSError("injected late verification failure"),
        ):
            with self.assertRaises(ExportError):
                export_handoff(project, root, now=clock, export_id="atomic")
        self.assertFalse((root / "Synthetic-Civil_export_atomic").exists())
        self.assertFalse(
            any(
                path.name.startswith(
                    ".Synthetic-Civil_export_atomic.staging-"
                )
                for path in root.iterdir()
            )
        )
        self.assertEqual(project.export_history, prior_history)
        self.assertEqual(project.exports_stale, prior_stale)

    def test_en_coordinate_preset_is_explicit(self):
        root = fresh_dir()
        project, flow = project_and_workflow()
        point = add_approved(flow)
        result = export_handoff(
            project, root, now=clock, export_id="en", coordinate_order="EN"
        )
        with (result["export_dir"] / "Existing_Points.csv").open(
            encoding="utf-8-sig", newline=""
        ) as handle:
            row = next(csv.DictReader(handle))
        self.assertEqual(float(row["Northing"]), point.local_east)
        self.assertEqual(float(row["Easting"]), point.local_north)

    def test_missing_calibration_blocks_export(self):
        project, _flow = project_and_workflow(calibrated=False)
        with self.assertRaises(ExportError):
            export_handoff(project, fresh_dir(), now=clock, export_id="blocked")
