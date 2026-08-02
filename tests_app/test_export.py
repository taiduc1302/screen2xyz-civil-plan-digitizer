from __future__ import annotations

import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from screen2xyz_app.capture import CapturedPoint
from screen2xyz_app.export import POINT_HEADERS, export_csv, export_xlsx
from screen2xyz_app.mapping import ChannelMapping, ChannelSource
from screen2xyz_app.store import SessionStore


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.store = SessionStore(self.root)
        mapping = ChannelMapping({
            "x": ChannelSource("screen_zone_ocr", (0, 0, 20, 20)),
            "y": ChannelSource("screen_zone_ocr", (20, 0, 20, 20)),
            "z": ChannelSource("plan_label_ocr"),
        })
        self.session_id = self.store.start_session(mapping)
        self.store.append_point(self.session_id, CapturedPoint(
            values={"x": 1.25, "y": 2.5, "z": 49.78, "description": "=unsafe"},
            source_methods={"x": "screen_zone_ocr", "y": "screen_zone_ocr", "z": "plan_label_ocr"},
            raw_texts={"x": "1.25", "y": "2.5", "z": "49.78"},
            confidences={"x": 0.9, "y": 0.8, "z": 0.7},
        ))

    def tearDown(self):
        self.store.close()
        self.temporary.cleanup()

    def test_xlsx_has_points_and_session_sheets(self):
        path = export_xlsx(self.store, self.session_id, self.root / "points.xlsx")
        workbook = load_workbook(path, data_only=False)
        self.assertEqual(workbook.sheetnames, ["Points", "Session"])
        points = workbook["Points"]
        self.assertEqual(tuple(cell.value for cell in points[1]), POINT_HEADERS)
        self.assertEqual(points["B2"].value, 1.25)
        self.assertEqual(points["E2"].value, "'=unsafe")
        self.assertIn("conceptual estimating", workbook["Session"]["B2"].value)

    def test_csv_uses_same_columns_and_escapes_formula_text(self):
        path = export_csv(self.store, self.session_id, self.root / "points.csv")
        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
        self.assertEqual(tuple(rows[0]), POINT_HEADERS)
        self.assertEqual(rows[0]["Description"], "'=unsafe")


if __name__ == "__main__":
    unittest.main()
