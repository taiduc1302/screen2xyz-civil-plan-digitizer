"""Run the realistic image-to-XLSX proof and enforce its release gate."""

from __future__ import annotations

import argparse
import collections
import json
import shutil
import time
from dataclasses import asdict, dataclass
from pathlib import Path

from openpyxl import load_workbook

from screen2xyz_app.backends import OcrPolicy, ScreenOcrBackend
from screen2xyz_app.capture import AutoCaptureEngine, CapturePipeline
from screen2xyz_app.export import export_xlsx
from screen2xyz_app.mapping import ChannelMapping, ChannelSource
from screen2xyz_app.store import SessionStore
from screen2xyz_civil.ocr import TesseractOcrAdapter

from .renderers import PlanLabel, StatusFrame, render_plan, render_status_scenarios


@dataclass(frozen=True)
class HarnessMetrics:
    ground_truth_rows: int
    captured_rows: int
    correctly_captured_rows: int
    accuracy: float
    false_duplicate_rows: int
    hallucinated_rows: int
    parse_or_ocr_failures: int
    plan_labels_rendered: int
    plan_labels_used: int
    sqlite_xlsx_exact: bool
    duration_seconds: float
    tesseract: str

    @property
    def passed(self) -> bool:
        return (
            self.accuracy >= 0.95
            and self.false_duplicate_rows == 0
            and self.hallucinated_rows == 0
            and self.sqlite_xlsx_exact
        )


class RealFrameReader:
    """Backend adapter selecting real rendered bytes for the current tick."""

    def __init__(
        self,
        backend: ScreenOcrBackend,
        plan_path: Path,
        labels: list[PlanLabel],
    ) -> None:
        self.backend = backend
        self.plan_path = plan_path
        self.labels = labels
        self.frame: StatusFrame | None = None
        self.label_index = 0

    def select(self, frame: StatusFrame, label_index: int) -> None:
        self.frame = frame
        self.label_index = label_index

    def __call__(self, column, source, context):
        del context
        assert self.frame is not None
        if column in {"x", "y"}:
            zone = self.frame.x_zone if column == "x" else self.frame.y_zone
            ranges = {
                "point-light": {"x": (0.0, 1_000.0), "y": (-1_000.0, 0.0)},
                "point-grouped": {"x": (400_000.0, 500_000.0), "y": (5_000_000.0, 6_000_000.0)},
                "comma-dark": {"x": (10_000.0, 30_000.0), "y": (5_000.0, 20_000.0)},
                "comma-grouped": {"x": (400_000.0, 500_000.0), "y": (5_000_000.0, 6_000_000.0)},
            }
            return self.backend.read_image_region(
                self.frame.path,
                zone,
                cache_key=(column, self.frame.style),
                policy=OcrPolicy(
                    separator_mode=source.decimal_separator,
                    numeric_range=ranges[self.frame.style][column],
                    precision_min=2,
                    confidence_min=0.35,
                    psm_modes=(7, 8, 13),
                    upscale=4,
                ),
            )
        label = self.labels[self.label_index]
        return self.backend.read_image_region(
            self.plan_path,
            label.ocr_region,
            cache_key=("plan", label.label_id),
            policy=OcrPolicy(
                separator_mode="point",
                numeric_range=(30.0, 100.0),
                precision_min=2,
                consensus_min=2,
                confidence_min=0.25,
                psm_modes=(6, 7),
                rotation_angles=(0, -15, 15, -20, 20, -25, 25),
                upscale=2,
            ),
        )


def _tuple(row) -> tuple[float, float, float]:
    return (round(float(row["x"]), 6), round(float(row["y"]), 6), round(float(row["z"]), 6))


def _verify_xlsx(store: SessionStore, session_id: str, path: Path) -> bool:
    database_rows = store.points(session_id)
    workbook = load_workbook(path, data_only=False, read_only=True)
    sheet = workbook["Points"]
    exported = list(sheet.iter_rows(min_row=2, values_only=True))
    if len(exported) != len(database_rows):
        return False
    for sequence, (db, xlsx) in enumerate(zip(database_rows, exported), start=1):
        confidences = [
            db[key]
            for key in ("confidence_x", "confidence_y", "confidence_z")
            if db[key] is not None
        ]
        expected = (
            db["point_number"] or str(sequence),
            db["x"], db["y"], db["z"], db["description"] or None,
            db["source_method_x"], db["source_method_y"], db["source_method_z"],
            None if not confidences else sum(confidences) / len(confidences),
            db["created_utc"],
        )
        if tuple(xlsx) != expected:
            return False
    return True


def run_harness(output_dir: Path, *, count_per_style: int = 55) -> HarnessMetrics:
    started = time.perf_counter()
    executable = TesseractOcrAdapter.find_executable()
    if executable is None:
        raise RuntimeError("realistic harness requires a local Tesseract executable")
    output_dir = output_dir.resolve()
    if output_dir.exists():
        shutil.rmtree(output_dir)
    output_dir.mkdir(parents=True)
    scenarios = render_status_scenarios(output_dir / "status_frames", count_per_style)
    plan_path = output_dir / "fake_plan" / "synthetic_grading_plan.png"
    labels = render_plan(plan_path)
    backend = ScreenOcrBackend(executable=executable)
    reader = RealFrameReader(backend, plan_path, labels)

    expected: list[tuple[float, float, float]] = []
    captured: list[tuple[float, float, float]] = []
    failures: list[dict[str, object]] = []
    xlsx_exact = True
    project_dir = output_dir / "project"
    store = SessionStore(project_dir)
    global_index = 0
    try:
        for style, primer, frames in scenarios:
            mapping = ChannelMapping({
                "x": ChannelSource("screen_zone_ocr", primer.x_zone, style.decimal_separator),
                "y": ChannelSource("screen_zone_ocr", primer.y_zone, style.decimal_separator),
                "z": ChannelSource("screen_zone_ocr", (0, 0, 120, 60), "point"),
            })
            session_id = store.start_session(mapping)
            engine = AutoCaptureEngine(
                CapturePipeline(mapping, reader),
                lambda point, sid=session_id: store.append_point(sid, point),
            )
            reader.select(primer, 0)
            try:
                engine.poll()
                engine.poll()
            except (ValueError, RuntimeError) as exc:
                raise RuntimeError(f"primer OCR failed for {style.name}: {exc}") from exc
            for frame in frames:
                label_index = global_index % len(labels)
                label = labels[label_index]
                expected.append((frame.x, frame.y, label.value))
                reader.select(frame, label_index)
                try:
                    engine.poll()
                    engine.poll()
                except (ValueError, RuntimeError) as exc:
                    failures.append({
                        "sequence": frame.sequence,
                        "style": style.name,
                        "label": label.label_id,
                        "error": str(exc),
                    })
                global_index += 1
            session_rows = store.points(session_id)
            captured.extend(_tuple(row) for row in session_rows)
            export_path = export_xlsx(
                store, session_id, output_dir / "exports" / f"{style.name}.xlsx"
            )
            xlsx_exact = xlsx_exact and _verify_xlsx(store, session_id, export_path)
    finally:
        store.close()

    expected_by_xy = {(x, y): z for x, y, z in expected}
    captured_xy = collections.Counter((x, y) for x, y, _z in captured)
    correct = sum(
        1 for x, y, z in captured if expected_by_xy.get((x, y)) == z
    )
    false_duplicates = sum(
        max(0, count - 1)
        for xy, count in captured_xy.items()
        if xy in expected_by_xy
    )
    hallucinated = sum(
        1 for x, y, z in captured if expected_by_xy.get((x, y)) != z
    )
    metrics = HarnessMetrics(
        ground_truth_rows=len(expected),
        captured_rows=len(captured),
        correctly_captured_rows=correct,
        accuracy=(correct / len(expected) if expected else 0.0),
        false_duplicate_rows=false_duplicates,
        hallucinated_rows=hallucinated,
        parse_or_ocr_failures=len(failures),
        plan_labels_rendered=len(labels),
        plan_labels_used=len({index % len(labels) for index in range(len(expected))}),
        sqlite_xlsx_exact=xlsx_exact,
        duration_seconds=round(time.perf_counter() - started, 3),
        tesseract=str(executable),
    )
    report = {
        "schema_version": "1.0",
        "metrics": asdict(metrics),
        "passed": metrics.passed,
        "thresholds": {
            "accuracy_min": 0.95,
            "false_duplicate_rows_max": 0,
            "hallucinated_rows_max": 0,
            "sqlite_xlsx_exact": True,
        },
        "failures": failures,
    }
    (output_dir / "metrics.json").write_text(
        json.dumps(report, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    (output_dir / "summary.md").write_text(
        "# Screen2XYZ realistic OCR harness\n\n"
        f"- Accuracy: {metrics.accuracy:.2%} "
        f"({metrics.correctly_captured_rows}/{metrics.ground_truth_rows})\n"
        f"- False duplicates: {metrics.false_duplicate_rows}\n"
        f"- Hallucinated rows: {metrics.hallucinated_rows}\n"
        f"- OCR/parse failures: {metrics.parse_or_ocr_failures}\n"
        f"- SQLite/XLSX exact: {metrics.sqlite_xlsx_exact}\n"
        f"- Result: {'PASS' if metrics.passed else 'FAIL'}\n",
        encoding="utf-8",
    )
    return metrics


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=Path(".lab_work/v2_5_harness"))
    parser.add_argument("--count-per-style", type=int, default=55)
    args = parser.parse_args(argv)
    metrics = run_harness(args.output, count_per_style=args.count_per_style)
    print(json.dumps(asdict(metrics), indent=2, sort_keys=True))
    return 0 if metrics.passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
