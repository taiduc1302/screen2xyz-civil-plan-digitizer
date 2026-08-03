from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

from screen2xyz_app.backends import (
    CursorOcrCandidate,
    DefaultReader,
    OcrPolicy,
    ScreenOcrBackend,
)
from screen2xyz_app.controller import CaptureSessionController
from screen2xyz_app.mapping import ChannelMapping, ChannelSource
from screen2xyz_civil.detection import Rect


class CursorCaptureTests(unittest.TestCase):
    def _backend(self, candidates, grabbed):
        def grab(zone):
            grabbed.append(zone)
            return Image.new("RGB", (zone[2], zone[3]), "white")

        return ScreenOcrBackend(
            executable=Path("tesseract"),
            capture_provider=grab,
            cursor_candidate_provider=lambda _image, _policy: tuple(candidates),
        )

    def test_cursor_box_is_centred_and_nearest_candidate_wins(self):
        grabbed = []
        candidates = (
            CursorOcrCandidate("49.47", 49.47, 0.91, Rect(8, 20, 38, 40), 0),
            CursorOcrCandidate("49.78", 49.78, 0.94, Rect(66, 20, 96, 40), 0),
        )
        reader = DefaultReader(
            self._backend(candidates, grabbed),
            cursor_position_provider=lambda: (500, 400),
        )
        source = ChannelSource(
            "screen_cursor_ocr",
            cursor_box_size=(160, 60),
            numeric_range=(40.0, 60.0),
        )
        reading = reader("z", source, {})
        self.assertEqual(grabbed, [(420, 370, 160, 60)])
        self.assertEqual(reading.raw_text, "49.78")
        self.assertAlmostEqual(reading.confidence, 0.94)

    def test_cursor_candidate_outside_snap_radius_is_rejected(self):
        grabbed = []
        candidates = (
            CursorOcrCandidate("49.47", 49.47, 0.91, Rect(0, 0, 10, 10), 0),
        )
        reader = DefaultReader(
            self._backend(candidates, grabbed),
            cursor_position_provider=lambda: (500, 400),
        )
        source = ChannelSource("screen_cursor_ocr", cursor_box_size=(160, 60))
        with self.assertRaisesRegex(ValueError, "near the cursor"):
            reader("z", source, {})

    def test_rotated_candidate_uses_multi_angle_policy(self):
        seen_angles = []

        def provider(_image, policy: OcrPolicy):
            seen_angles.extend(policy.rotation_angles)
            return (
                CursorOcrCandidate("49.78", 49.78, 0.92, Rect(68, 20, 96, 40), 20),
            )

        backend = ScreenOcrBackend(
            executable=Path("tesseract"),
            capture_provider=lambda zone: Image.new("RGB", (zone[2], zone[3]), "white"),
            cursor_candidate_provider=provider,
        )
        reader = DefaultReader(backend, cursor_position_provider=lambda: (100, 100))
        reading = reader(
            "z", ChannelSource("screen_cursor_ocr", cursor_box_size=(160, 60)), {}
        )
        self.assertEqual(reading.raw_text, "49.78")
        self.assertIn(20, seen_angles)
        self.assertIn(-20, seen_angles)

    def test_mixed_cursor_session_lands_in_sqlite_and_xlsx(self):
        class MixedBackend:
            def read_zone(self, zone, *, policy):
                value = "2380658" if zone[0] == 10 else "1328135"
                from screen2xyz_app.capture import Reading
                return Reading(value, 0.95)

            def read_cursor(self, cursor, box_size, *, policy, snap_radius_px):
                del cursor, box_size, policy, snap_radius_px
                from screen2xyz_app.capture import Reading
                return Reading("49.78", 0.93)

        mapping = ChannelMapping({
            "x": ChannelSource("screen_zone_ocr", (10, 10, 120, 28)),
            "y": ChannelSource("screen_zone_ocr", (140, 10, 120, 28)),
            "z": ChannelSource("screen_cursor_ocr", cursor_box_size=(160, 60)),
        })
        self.assertTrue(mapping.automatic)
        with tempfile.TemporaryDirectory() as temporary:
            controller = CaptureSessionController(
                Path(temporary),
                mapping,
                reader=DefaultReader(
                    MixedBackend(), cursor_position_provider=lambda: (600, 500)
                ),
            )
            controller.capture_click()
            rows = controller.points()
            self.assertEqual(len(rows), 1)
            self.assertEqual((rows[0]["x"], rows[0]["y"], rows[0]["z"]),
                             (2380658.0, 1328135.0, 49.78))
            output = controller.export_xlsx(Path(temporary) / "mixed.xlsx")
            sheet = load_workbook(output, data_only=True)["Points"]
            exported = list(sheet.iter_rows(min_row=2, values_only=True))
            self.assertEqual(exported[0][1:4], (2380658.0, 1328135.0, 49.78))
            controller.close()


if __name__ == "__main__":
    unittest.main()
